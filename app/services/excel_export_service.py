import io
import logging
import re
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.config import settings
from app.models.candidate import Candidate
from app.models.client import Client
from app.models.intelligence import ClientIntelligence
from app.models.job import Job
from app.services.geo_resolver_service import geo_resolver

logger = logging.getLogger("app.services.excel")


# Palette Definitions
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
ZONE_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
WHITE_BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")

# Functional Zone Header Fills
ZONE_FILLS = {
    "Zone 1: Deal Economics": PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid"),        # Dark Navy
    "Zone 2: Competition": PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid"),           # Slate Blue
    "Zone 3: Client Credibility": PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid"),    # Deep Steel Blue
    "Zone 4: Extracted Intelligence": PatternFill(start_color="203764", end_color="203764", fill_type="solid"),# Royal Blue
    "Zone 5: Decision-Maker Outreach": PatternFill(start_color="134F5C", end_color="134F5C", fill_type="solid"),# Dark Teal
    "Zone 6: Sales Hooks & Scope": PatternFill(start_color="4A235A", end_color="4A235A", fill_type="solid"),   # Deep Purple
}

# Tier Styling
TIER_A_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Soft Green
TIER_A_FONT = Font(name="Calibri", size=10, bold=True, color="276A3C")

TIER_B_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Soft Yellow
TIER_B_FONT = Font(name="Calibri", size=10, bold=True, color="8C6B00")

TIER_C_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")  # Soft Gray
TIER_C_FONT = Font(name="Calibri", size=10, color="7F7F7F")

# Borders
THIN_BORDER_SIDE = Side(border_style="thin", color="D9D9D9")
GRID_BORDER = Border(
    left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE
)


# Column Definitions (46 Master Columns in 6 Functional Zones)
COLUMNS_SPEC = [
    # Zone 1: Opportunity & Deal Economics (Cols 1–9)
    {"id": "job_title", "title": "Job Title", "zone": "Zone 1: Deal Economics", "width": 32, "type": "text"},
    {"id": "job_url", "title": "Job URL", "zone": "Zone 1: Deal Economics", "width": 46, "type": "url"},
    {"id": "posted_at_relative", "title": "Posted Relative", "zone": "Zone 1: Deal Economics", "width": 16, "type": "text"},
    {"id": "posted_at_iso", "title": "Posted ISO", "zone": "Zone 1: Deal Economics", "width": 22, "type": "text"},
    {"id": "pay_type", "title": "Pay Type", "zone": "Zone 1: Deal Economics", "width": 14, "type": "text"},
    {"id": "budget_min_usd", "title": "Budget Min ($)", "zone": "Zone 1: Deal Economics", "width": 15, "type": "currency"},
    {"id": "budget_max_usd", "title": "Budget Max ($)", "zone": "Zone 1: Deal Economics", "width": 15, "type": "currency"},
    {"id": "project_duration", "title": "Duration", "zone": "Zone 1: Deal Economics", "width": 18, "type": "text"},
    {"id": "experience_level", "title": "Experience Level", "zone": "Zone 1: Deal Economics", "width": 16, "type": "text"},

    # Zone 2: Competition & Outreach Urgency (Cols 10–14)
    {"id": "proposals_tier", "title": "Proposals Tier", "zone": "Zone 2: Competition", "width": 16, "type": "text"},
    {"id": "interviewing_count", "title": "Interviewing", "zone": "Zone 2: Competition", "width": 14, "type": "int"},
    {"id": "invites_sent_count", "title": "Invites Sent", "zone": "Zone 2: Competition", "width": 14, "type": "int"},
    {"id": "unanswered_invites", "title": "Unanswered Invites", "zone": "Zone 2: Competition", "width": 18, "type": "int"},
    {"id": "last_client_activity", "title": "Client Local Time", "zone": "Zone 2: Competition", "width": 20, "type": "text"},

    # Zone 3: Client Credibility & Financial Power (Cols 15–26)
    {"id": "payment_verified", "title": "Payment Verified", "zone": "Zone 3: Client Credibility", "width": 16, "type": "bool"},
    {"id": "client_rating", "title": "Client Rating", "zone": "Zone 3: Client Credibility", "width": 14, "type": "decimal"},
    {"id": "client_reviews_count", "title": "Total Reviews", "zone": "Zone 3: Client Credibility", "width": 14, "type": "int"},
    {"id": "client_total_spent_usd", "title": "Total Spent ($)", "zone": "Zone 3: Client Credibility", "width": 16, "type": "currency"},
    {"id": "client_hire_rate_pct", "title": "Hire Rate (%)", "zone": "Zone 3: Client Credibility", "width": 14, "type": "pct"},
    {"id": "client_total_hires", "title": "Total Hires", "zone": "Zone 3: Client Credibility", "width": 13, "type": "int"},
    {"id": "client_active_hires", "title": "Active Hires", "zone": "Zone 3: Client Credibility", "width": 13, "type": "int"},
    {"id": "client_avg_hourly_rate", "title": "Avg Hourly ($)", "zone": "Zone 3: Client Credibility", "width": 15, "type": "currency"},
    {"id": "client_city", "title": "Client City", "zone": "Zone 3: Client Credibility", "width": 16, "type": "text"},
    {"id": "client_country", "title": "Client Country", "zone": "Zone 3: Client Credibility", "width": 16, "type": "text"},
    {"id": "client_local_time", "title": "Timezone / Local Time", "zone": "Zone 3: Client Credibility", "width": 22, "type": "text"},
    {"id": "client_member_since", "title": "Member Since", "zone": "Zone 3: Client Credibility", "width": 16, "type": "text"},

    # Zone 4: Extracted Client & Company Intelligence (Cols 27–35)
    {"id": "client_first_name", "title": "Client Name", "zone": "Zone 4: Extracted Intelligence", "width": 18, "type": "text"},
    {"id": "client_name_confidence", "title": "Name Conf (%)", "zone": "Zone 4: Extracted Intelligence", "width": 14, "type": "pct"},
    {"id": "client_name_mentions", "title": "Name Mentions", "zone": "Zone 4: Extracted Intelligence", "width": 15, "type": "int"},
    {"id": "client_name_evidence", "title": "Name Evidence Snippet", "zone": "Zone 4: Extracted Intelligence", "width": 38, "type": "text"},
    {"id": "company_name", "title": "Company Name", "zone": "Zone 4: Extracted Intelligence", "width": 20, "type": "text"},
    {"id": "company_evidence", "title": "Company Evidence Snippet", "zone": "Zone 4: Extracted Intelligence", "width": 38, "type": "text"},
    {"id": "is_agency", "title": "Is Agency", "zone": "Zone 4: Extracted Intelligence", "width": 13, "type": "bool"},
    {"id": "agency_detection_reason", "title": "Agency Reason", "zone": "Zone 4: Extracted Intelligence", "width": 28, "type": "text"},
    {"id": "project_brand_keywords", "title": "Project Keywords", "zone": "Zone 4: Extracted Intelligence", "width": 26, "type": "text"},

    # Zone 5: Direct LinkedIn Decision-Maker Outreach (Cols 36–43)
    {"id": "decision_maker_name", "title": "Decision Maker Name", "zone": "Zone 5: Decision-Maker Outreach", "width": 22, "type": "text"},
    {"id": "decision_maker_title", "title": "Headline / Role", "zone": "Zone 5: Decision-Maker Outreach", "width": 35, "type": "text"},
    {"id": "decision_maker_location", "title": "Candidate Location", "zone": "Zone 5: Decision-Maker Outreach", "width": 26, "type": "text"},
    {"id": "match_score_pct", "title": "Match Score (%)", "zone": "Zone 5: Decision-Maker Outreach", "width": 15, "type": "pct"},
    {"id": "match_tier", "title": "Match Tier", "zone": "Zone 5: Decision-Maker Outreach", "width": 20, "type": "tier"},
    {"id": "linkedin_profile_url", "title": "LinkedIn Profile URL", "zone": "Zone 5: Decision-Maker Outreach", "width": 50, "type": "url"},
    {"id": "linkedin_search_url_metro", "title": "LinkedIn Search (City GeoURN)", "zone": "Zone 5: Decision-Maker Outreach", "width": 60, "type": "url"},
    {"id": "linkedin_search_url_country", "title": "LinkedIn Search (Country GeoURN)", "zone": "Zone 5: Decision-Maker Outreach", "width": 60, "type": "url"},

    # Zone 6: Sales Pitch Hooks & Technical Scope (Cols 44–46)
    {"id": "skills_required", "title": "Skills Required", "zone": "Zone 6: Sales Hooks & Scope", "width": 32, "type": "text"},
    {"id": "job_description", "title": "Full Job Description", "zone": "Zone 6: Sales Hooks & Scope", "width": 55, "type": "text"},
    {"id": "recommended_pitch_hook", "title": "Recommended Pitch Hook", "zone": "Zone 6: Sales Hooks & Scope", "width": 55, "type": "text"},
]


def generate_sales_pitch_hook(
    client_name: Optional[str],
    company_name: Optional[str],
    job_title: str,
    skills_tags: Optional[str],
) -> str:
    """Generates an executive-ready opening sales pitch hook tailored to client, company, and project."""
    top_skills = [s.strip() for s in (skills_tags or "").split(",") if s.strip()][:3]
    skills_phrase = f"in {', '.join(top_skills)}" if top_skills else "in your required technology stack"

    if client_name and client_name != "Client" and company_name:
        return (
            f"Hi {client_name}, saw your {job_title} opening for {company_name}. "
            f"With proven production experience {skills_phrase}, I can deliver immediate impact and streamline execution."
        )
    elif client_name and client_name != "Client":
        return (
            f"Hi {client_name}, I came across your {job_title} posting. "
            f"Having delivered multiple high-reliability solutions {skills_phrase}, I can help you hit the ground running."
        )
    elif company_name:
        return (
            f"Hi there, saw your {job_title} project for {company_name}. "
            f"With strong expertise {skills_phrase}, I would welcome the opportunity to support your technical roadmap."
        )
    return (
        f"Hi, I saw your {job_title} posting. "
        f"With hands-on expertise {skills_phrase}, I can deliver high-quality, maintainable results on schedule."
    )


class ExcelExportService:
    """Enterprise Master Excel Exporter with 6 Functional Zones, Raw Clickable URLs, and Tier Badging."""

    @staticmethod
    def build_row_data(
        job: Job,
        candidate: Optional[Candidate] = None,
    ) -> Dict[str, Any]:
        """Maps Job, Client, Intelligence, and Candidate models into the 46-column schema dictionary."""
        client: Optional[Client] = job.client
        intel: Optional[ClientIntelligence] = job.intelligence

        # In-app search links via GeoResolver
        client_name = (intel.confirmed_client_name if intel and intel.confirmed_client_name else "Client")
        city = client.city if client else None
        country = client.country if client else None

        metro_res = geo_resolver.resolve(city=city, country=country)
        metro_urn = metro_res["geo_urn"] if metro_res and metro_res["level"] == "metro" else None
        country_res = geo_resolver.resolve(country=country)
        country_urn = country_res["geo_urn"] if country_res else None

        url_metro = (
            geo_resolver.build_inapp_search_url(keywords=client_name, geo_urn=metro_urn)
            if metro_urn
            else geo_resolver.build_inapp_search_url(keywords=client_name)
        )
        url_country = (
            geo_resolver.build_inapp_search_url(keywords=client_name, geo_urn=country_urn)
            if country_urn
            else geo_resolver.build_inapp_search_url(keywords=client_name)
        )

        pitch_hook = generate_sales_pitch_hook(
            client_name=client_name,
            company_name=intel.confirmed_company_name if intel else None,
            job_title=job.title,
            skills_tags=job.skills_tags,
        )

        return {
            # Zone 1
            "job_title": job.title,
            "job_url": job.job_url,
            "posted_at_relative": "Recently",
            "posted_at_iso": job.posted_at.isoformat() if job.posted_at else "",
            "pay_type": job.pay_type,
            "budget_min_usd": float(job.budget_min_usd) if job.budget_min_usd is not None else None,
            "budget_max_usd": float(job.budget_max_usd) if job.budget_max_usd is not None else None,
            "project_duration": job.project_duration or "Not specified",
            "experience_level": job.experience_level or "Not specified",

            # Zone 2
            "proposals_tier": job.proposals_tier or "Less than 5",
            "interviewing_count": job.interviewing_count,
            "invites_sent_count": job.invites_sent_count,
            "unanswered_invites": job.unanswered_invites,
            "last_client_activity": client.local_time if client and client.local_time else "Active recently",

            # Zone 3
            "payment_verified": client.payment_verified if client else False,
            "client_rating": float(client.client_rating) if client and client.client_rating else None,
            "client_reviews_count": client.total_reviews if client else 0,
            "client_total_spent_usd": float(client.total_spent_usd) if client else 0.0,
            "client_hire_rate_pct": float(client.hire_rate_pct) if client and client.hire_rate_pct else None,
            "client_total_hires": client.total_hires if client else 0,
            "client_active_hires": client.active_hires if client else 0,
            "client_avg_hourly_rate": float(client.avg_hourly_rate_usd) if client and client.avg_hourly_rate_usd else None,
            "client_city": client.city if client else None,
            "client_country": client.country if client else None,
            "client_local_time": client.local_time if client else None,
            "client_member_since": client.member_since if client else None,

            # Zone 4
            "client_first_name": intel.confirmed_client_name if intel else None,
            "client_name_confidence": float(intel.name_confidence) if intel else None,
            "client_name_mentions": intel.name_mentions_count if intel else 0,
            "client_name_evidence": intel.name_evidence_snippet if intel else None,
            "company_name": intel.confirmed_company_name if intel else None,
            "company_evidence": intel.company_evidence_snippet if intel else None,
            "is_agency": intel.is_agency if intel else False,
            "agency_detection_reason": intel.agency_detection_reason if intel else None,
            "project_brand_keywords": job.skills_tags or intel.confirmed_company_name if intel else None,

            # Zone 5
            "decision_maker_name": candidate.full_name if candidate else "Search Link Only",
            "decision_maker_title": candidate.headline if candidate else "Pending Resolution",
            "decision_maker_location": candidate.profile_location if candidate else (f"{city}, {country}" if city else country),
            "match_score_pct": float(candidate.match_score) if candidate else None,
            "match_tier": candidate.tier if candidate else "Search Links Only",
            "linkedin_profile_url": candidate.linkedin_url if candidate else url_metro,
            "linkedin_search_url_metro": url_metro,
            "linkedin_search_url_country": url_country,

            # Zone 6
            "skills_required": job.skills_tags,
            "job_description": job.description,
            "recommended_pitch_hook": pitch_hook,
        }

    def generate_leads_workbook(
        self,
        jobs: List[Job],
        output_path: Optional[Path] = None,
        only_primary: bool = False,
    ) -> Path:
        """
        Creates an executive-ready Excel workbook with 46 master columns across 6 functional zones.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Decision Makers & Leads"
        ws.views.sheetView[0].showGridLines = True

        # 1. Title Block (Row 1)
        ws.merge_cells("A1:AT1")
        title_cell = ws["A1"]
        title_cell.value = "Upwork Client & LinkedIn Decision-Maker Master Intelligence Export"
        title_cell.font = TITLE_FONT
        title_cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[1].height = 34

        # 2. Subtitle / Metadata Block (Row 2)
        ws.merge_cells("A2:AT2")
        sub_cell = ws["A2"]
        sub_cell.value = (
            f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} | "
            f"Total Jobs Analyzed: {len(jobs)} | "
            "Clickable LinkedIn & Upwork URLs | 6 Functional Zones"
        )
        sub_cell.font = SUBTITLE_FONT
        sub_cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[2].height = 20

        # 3. Zone Group Header (Row 3)
        ws.row_dimensions[3].height = 22
        zone_ranges = [
            ("A3:I3", "Zone 1: Deal Economics (Cols A–I)", ZONE_FILLS["Zone 1: Deal Economics"]),
            ("J3:N3", "Zone 2: Competition (Cols J–N)", ZONE_FILLS["Zone 2: Competition"]),
            ("O3:Z3", "Zone 3: Client Credibility (Cols O–Z)", ZONE_FILLS["Zone 3: Client Credibility"]),
            ("AA3:AI3", "Zone 4: Extracted Intelligence (Cols AA–AI)", ZONE_FILLS["Zone 4: Extracted Intelligence"]),
            ("AJ3:AQ3", "Zone 5: Decision-Maker Outreach (Cols AJ–AQ)", ZONE_FILLS["Zone 5: Decision-Maker Outreach"]),
            ("AR3:AT3", "Zone 6: Sales Hooks & Scope (Cols AR–AT)", ZONE_FILLS["Zone 6: Sales Hooks & Scope"]),
        ]
        for rng, label, fill in zone_ranges:
            ws.merge_cells(rng)
            first_cell = ws[rng.split(":")[0]]
            first_cell.value = label
            first_cell.font = ZONE_HEADER_FONT
            first_cell.fill = fill
            first_cell.alignment = Alignment(horizontal="center", vertical="center")

        # 4. Column Headers (Row 4)
        ws.row_dimensions[4].height = 28
        for col_idx, col_spec in enumerate(COLUMNS_SPEC, 1):
            cell = ws.cell(row=4, column=col_idx, value=col_spec["title"])
            cell.fill = ZONE_FILLS.get(col_spec["zone"], ZONE_FILLS["Zone 1: Deal Economics"])
            cell.font = WHITE_BOLD_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = GRID_BORDER

        # 5. Data Rows (Row 5+)
        current_row = 5
        for job in jobs:
            cands = job.candidates or []
            if only_primary:
                cands = [c for c in cands if c.is_primary_resolved]
                if not cands and job.candidates:
                    cands = [job.candidates[0]]

            # If no candidates discovered yet, still output 1 row with job & intelligence
            records_to_export = cands if cands else [None]

            for cand in records_to_export:
                row_dict = self.build_row_data(job, cand)
                ws.row_dimensions[current_row].height = 26

                tier_val = str(row_dict.get("match_tier") or "")

                for col_idx, col_spec in enumerate(COLUMNS_SPEC, 1):
                    val = row_dict.get(col_spec["id"])
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.border = GRID_BORDER

                    col_type = col_spec["type"]
                    align = Alignment(vertical="center")

                    # Tier Column Styling
                    if col_type == "tier":
                        align = Alignment(horizontal="center", vertical="center")
                        cell.value = str(val or "")
                        if "Tier A" in tier_val:
                            cell.fill, cell.font = TIER_A_FILL, TIER_A_FONT
                        elif "Tier B" in tier_val:
                            cell.fill, cell.font = TIER_B_FILL, TIER_B_FONT
                        else:
                            cell.fill, cell.font = TIER_C_FILL, TIER_C_FONT

                    # Hyperlink Columns (Raw Clickable URLs)
                    elif col_type == "url":
                        url_str = str(val or "")
                        if url_str.startswith("http"):
                            cell.value = url_str
                            cell.hyperlink = url_str
                            cell.font = LINK_FONT
                        else:
                            cell.value = url_str
                            cell.font = REGULAR_FONT
                        align = Alignment(vertical="center")

                    # Percentage Columns
                    elif col_type == "pct":
                        align = Alignment(horizontal="center", vertical="center")
                        if val is not None:
                            cell.value = f"{float(val):.1f}%"
                            cell.font = BOLD_FONT
                        else:
                            cell.value = ""
                            cell.font = REGULAR_FONT

                    # Currency Columns
                    elif col_type == "currency":
                        align = Alignment(horizontal="right", vertical="center")
                        if val is not None:
                            cell.value = f"${float(val):,.2f}"
                            cell.font = REGULAR_FONT
                        else:
                            cell.value = ""
                            cell.font = REGULAR_FONT

                    # Integer / Decimal Columns
                    elif col_type in ("int", "decimal"):
                        align = Alignment(horizontal="center", vertical="center")
                        cell.value = val if val is not None else ""
                        cell.font = REGULAR_FONT

                    # Boolean Columns
                    elif col_type == "bool":
                        align = Alignment(horizontal="center", vertical="center")
                        cell.value = "TRUE" if val else "FALSE"
                        cell.font = BOLD_FONT if val else REGULAR_FONT

                    # Default Text
                    else:
                        cell.value = str(val) if val is not None else ""
                        cell.font = REGULAR_FONT

                    cell.alignment = align

                current_row += 1

        # 6. Auto-fit column widths using spec dimensions
        for col_idx, col_spec in enumerate(COLUMNS_SPEC, 1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = col_spec["width"]

        # 7. Safe File Saving (with Fallback for Locked Files)
        target_path = output_path or (settings.EXPORTS_DIR / "leads_decision_makers.xlsx")
        target_path.parent.mkdir(parents=True, exist_ok=True)

        try:
            wb.save(target_path)
            return target_path
        except PermissionError:
            fallback = target_path.parent / f"{target_path.stem}_new{target_path.suffix}"
            wb.save(fallback)
            return fallback

    async def export_from_db(
        self,
        session: AsyncSession,
        job_uid: Optional[str] = None,
        min_spend: Optional[Decimal] = None,
        tier: Optional[str] = None,
        country: Optional[str] = None,
        only_primary: bool = False,
        output_path: Optional[Path] = None,
    ) -> Path:
        """
        Dynamically filters database jobs and generates the Master 46-Column Excel Workbook.
        """
        stmt = (
            select(Job)
            .options(
                selectinload(Job.client),
                selectinload(Job.intelligence),
                selectinload(Job.candidates),
            )
            .order_by(Job.created_at.desc())
        )

        if job_uid:
            stmt = stmt.filter(Job.job_uid == job_uid)

        if country:
            stmt = stmt.join(Job.client).filter(Client.country.ilike(f"%{country}%"))

        if min_spend is not None:
            stmt = stmt.join(Job.client).filter(Client.total_spent_usd >= min_spend)

        res = await session.execute(stmt)
        jobs = res.scalars().all()

        # Optional filter by candidate tier
        if tier:
            filtered_jobs = []
            for j in jobs:
                has_tier = any(tier.lower() in (c.tier or "").lower() for c in j.candidates)
                if has_tier:
                    filtered_jobs.append(j)
            jobs = filtered_jobs

        logger.info(f"[EXCEL] Generating leads workbook for {len(jobs)} jobs (output: {output_path or 'default'})...")
        path = self.generate_leads_workbook(
            jobs=jobs,
            output_path=output_path,
            only_primary=only_primary,
        )
        logger.info(f"[EXCEL] Successfully generated sales-ready Excel workbook: {path}")
        return path


excel_export_service = ExcelExportService()
